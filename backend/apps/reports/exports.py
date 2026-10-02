"""Excel and CSV files for the report centre's tables."""
import csv
import io
from datetime import date

import openpyxl
from django.http import HttpResponse
from openpyxl.styles import Alignment, Font

from .views import (
    BRAND_BLUE, SUBHDR_FILL, THIN_BORDER, _col_width, _write_data_row, _write_header_row, _write_title,
)

XLSX = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
NUMERIC = ('number', 'kg', 'money', 'percent')
FORMATS = {'kg': '#,##0.00', 'money': '#,##0.00', 'percent': '0.0"%"'}
UNITS = {'kg': ' (kg)', 'money': ' (Rs.)', 'percent': ' (%)'}


def _value(value, kind):
    """A cell value: numbers as numbers so they can be summed in the sheet, dates as dates."""
    if value is None or value == '':
        return None
    if kind in NUMERIC:
        try:
            return float(value)
        except (TypeError, ValueError):
            return value
    if kind == 'date':
        try:
            return date.fromisoformat(str(value)[:10])
        except ValueError:
            return value
    return value


def _headers(report):
    return [c['label'] + UNITS.get(c['kind'], '') for c in report['columns']]


def _filename(report, extension):
    period = report['period']
    days = '_to_'.join(d for d in (period['start'], period['end']) if d) or 'all_time'
    return f"{report['key'].replace('-', '_')}_{days}.{extension}"


def _attachment(content, content_type, filename):
    response = HttpResponse(content, content_type=content_type)
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


def xlsx_response(report):
    columns = report['columns']
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = report['title'][:31]
    _write_title(ws, report['title'],
                 f"Period: {report['period']['label']} | Generated: {date.today().strftime('%d/%m/%Y')}",
                 max(len(columns), 2))
    _write_header_row(ws, 5, _headers(report))

    def write(row_number, values, alt=False):
        _write_data_row(ws, row_number, [_value(values.get(c['key']), c['kind']) for c in columns], alt=alt)
        for index, column in enumerate(columns, 1):
            cell = ws.cell(row_number, index)
            if column['kind'] in FORMATS:
                cell.number_format = FORMATS[column['kind']]
            if column['kind'] == 'date':
                cell.number_format = 'DD/MM/YYYY'
            if column['kind'] in NUMERIC:
                cell.alignment = Alignment(horizontal='right', vertical='center')

    row_number = 5
    for index, row in enumerate(report['rows'], 1):
        row_number = 5 + index
        write(row_number, row, alt=index % 2 == 0)
    if report['totals']:
        row_number += 1
        write(row_number, report['totals'])
        for index in range(1, len(columns) + 1):
            cell = ws.cell(row_number, index)
            cell.font = Font(bold=True, color='1E293B', size=10)
            cell.fill = SUBHDR_FILL
            cell.border = THIN_BORDER

    for index, column in enumerate(columns, 1):
        longest = max([len(str(row.get(column['key']) or '')) for row in report['rows']] + [len(column['label']) + 6])
        _col_width(ws, index, min(max(longest + 2, 12), 44))

    if report['notes']:
        row_number += 2
        ws.cell(row_number, 1, 'How these figures are worked out').font = Font(bold=True, color=BRAND_BLUE, size=11)
        for note in report['notes']:
            row_number += 1
            ws.cell(row_number, 1, note).font = Font(color='475569', size=9, italic=True)

    buffer = io.BytesIO()
    wb.save(buffer)
    return _attachment(buffer.getvalue(), XLSX, _filename(report, 'xlsx'))


def csv_response(report):
    columns = report['columns']
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(_headers(report))
    for row in report['rows'] + ([report['totals']] if report['totals'] else []):
        writer.writerow(['' if row.get(c['key']) is None else row.get(c['key']) for c in columns])
    # The BOM makes Excel read the file as UTF-8
    return _attachment('﻿' + buffer.getvalue(), 'text/csv; charset=utf-8', _filename(report, 'csv'))
